import docx
import os, sys
from typing import List
from pypdf import PdfReader

from langchain_core.documents import Document

def text_to_document(path: str) -> List[Document]:
    with open(path, "r", encoding="utf-8") as file:
        text = file.read()

    document = Document(
        page_content=text,
        metadata={"source": os.path.abspath(path)}
    )
    return [document]

def pdf_to_document(path: str) -> List[Document]:
    reader = PdfReader(path)
    documents = []

    for page_num, page in enumerate(reader.pages, start=1):
        text = page.extract_text()

        if not text.strip():
            continue

        document = Document(
            page_content=text,
            metadata={"source": f"{os.path.abspath(path)}: page: {page_num}: total_pages: {len(reader.pages)}"}
        )
        documents.append(document)

    return documents

def docx_to_document(path: str) -> List[Document]:
    doc_content = docx.Document(path)

    text = "\n\n".join([p.text for p in doc_content.paragraphs if p.text and p.text.strip()])

    document = Document(
        page_content=text,
        metadata={"source": os.path.abspath(path)}
    )
    return [document]

def load_documents(path: str) -> List[Document]:
    try:

        if not os.path.exists(path):
            raise FileNotFoundError(f"File not found on path: {os.path.abspath(path)}")

        else:
            ext = os.path.splitext(path)[1].lower()
            if ext in [".txt", ".md"]:
                return text_to_document(path=path)
            elif ext == ".pdf":
                return pdf_to_document(path=path)
            elif ext in [".docx", ".docs"]:
                return docx_to_document(path=path)

    except Exception as e:
        print(f"Got an Eception while trying to load the file: {e}")